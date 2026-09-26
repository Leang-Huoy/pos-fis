import openpyxl
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

xlsm_files = [f for f in os.listdir('.') if f.endswith('.xlsm') and not f.startswith('~$')]
print("Found files:", [f.encode('unicode_escape').decode() for f in xlsm_files])

target_file = xlsm_files[0]
print(f"Loading {target_file.encode('unicode_escape').decode()}...")

wb = openpyxl.load_workbook(target_file, data_only=True)
sheet_info = {}

for sheet in wb.sheetnames:
    ws = wb[sheet]
    rows_preview = []
    # inspect rows
    for r in range(1, min(ws.max_row + 1, 60)):
        row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else "" for c in range(1, min(ws.max_column + 1, 25))]
        if any(row_vals):
            rows_preview.append(f"R{r}: " + " | ".join(row_vals[:18]))
    
    sheet_info[sheet] = {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "preview_count": len(rows_preview),
        "rows": rows_preview[:40]
    }

with open("workbook_structure.txt", "w", encoding="utf-8") as f:
    for s_name, data in sheet_info.items():
        f.write(f"\n{'='*30} SHEET: {s_name} (rows: {data['max_row']}, cols: {data['max_column']}) {'='*30}\n")
        for r_line in data["rows"]:
            f.write(r_line + "\n")

print("Exported workbook structure to workbook_structure.txt successfully!")
